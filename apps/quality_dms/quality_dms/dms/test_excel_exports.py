"""Formatting tests for the three DMS Excel exports.

Every table cell must have a thin border on all four sides and the header row
must be bold, centered and wrapped. Banner rows above the table (the Assessment
Results Register title lines) are deliberately excluded from the border checks.
The data itself must be exactly what the on-screen report / table shows.
"""

import re
from io import BytesIO

import frappe
import openpyxl
from frappe.tests import IntegrationTestCase, UnitTestCase

from quality_dms.dms import api, excel_style
from quality_dms.dms.report.training_assessment_results_register import (
    training_assessment_results_register as assessment_register,
)
from quality_dms.dms.report.training_register import training_register

ASSESSMENT_SCORE_HEADERS = [
    "Sr. No.", "Assessment ID", "Training Title", "Date of completion",
    "Employee ID", "Employee Name", "Department", "Trainer",
    "Maximum Marks", "Marks Obtained", "Score (%)", "Pass Mark (%)",
    "Result (Pass/Fail)", "Remarks",
]


class TestExcelStyleHelper(UnitTestCase):
    def test_header_and_body_formats(self):
        styles = excel_style.bordered_styles(4, 3)
        header, body = styles["styles"][0], styles["styles"][1]
        self.assertEqual(body["border"], 1)
        self.assertEqual(header["border"], 1)
        self.assertEqual(header["bold"], 1)
        self.assertEqual(header["align"], "center")
        self.assertEqual(header["text_wrap"], 1)

    def test_every_table_cell_gets_a_style_including_blank_ones(self):
        styles = excel_style.bordered_styles(4, 3)
        self.assertEqual(sorted(styles["cell_styles"]), [(r, c) for r in range(3) for c in range(4)])
        self.assertTrue(all(styles["cell_styles"][(0, c)] == [0] for c in range(4)))
        self.assertTrue(all(styles["cell_styles"][(r, c)] == [1] for r in (1, 2) for c in range(4)))

    def test_no_column_wide_style_so_borders_stop_at_the_table(self):
        styles = excel_style.bordered_styles(4, 3)
        self.assertNotIn("column_styles", styles)
        self.assertNotIn("row_styles", styles)

    def test_extra_column_formats_layer_on_top_of_the_border_style(self):
        styles = excel_style.bordered_styles(3, 2, column_formats={1: {"bg_color": "#FFFF00"}})
        fill_id = len(styles["styles"]) - 1
        self.assertEqual(styles["styles"][fill_id], {"bg_color": "#FFFF00"})
        self.assertEqual(styles["cell_styles"][(0, 1)], [0, fill_id])
        self.assertEqual(styles["cell_styles"][(1, 1)], [1, fill_id])
        self.assertEqual(styles["cell_styles"][(1, 0)], [1])

    def test_column_widths_fit_content_within_limits(self):
        rows = [["Sr. No.", "Training Title"], [1, "x" * 200], [2, "short"]]
        widths = excel_style.column_widths(rows)
        self.assertEqual(len(widths), 2)
        self.assertGreaterEqual(widths[0], excel_style.MIN_WIDTH)
        self.assertEqual(widths[1], excel_style.MAX_WIDTH)

    def test_column_widths_ignore_empty_cells(self):
        self.assertEqual(len(excel_style.column_widths([["A", "B"], [None, ""]])), 2)

    def test_numbers_and_dates_get_the_center_style_when_rows_are_given(self):
        import datetime

        rows = [
            ["H1", "H2", "H3", "H4", "H5"],
            ["text", 5, 2.5, datetime.date(2026, 1, 1), True],
            ["007", None, "", datetime.datetime(2026, 1, 1, 10, 0), "x"],
        ]
        cells = excel_style.bordered_styles(5, 3, rows=rows)["cell_styles"]
        center = excel_style.CENTER_STYLE_ID
        self.assertEqual(cells[(1, 0)], [1])  # text
        self.assertEqual(cells[(1, 1)], [1, center])  # int
        self.assertEqual(cells[(1, 2)], [1, center])  # float
        self.assertEqual(cells[(1, 3)], [1, center])  # date
        self.assertEqual(cells[(1, 4)], [1])  # bool is not a number here
        self.assertEqual(cells[(2, 0)], [1])  # numeric-looking text stays text
        self.assertEqual(cells[(2, 1)], [1])  # None
        self.assertEqual(cells[(2, 2)], [1])  # ""
        self.assertEqual(cells[(2, 3)], [1, center])  # datetime
        self.assertEqual(cells[(0, 1)], [0])  # the header never takes the number style

    def test_center_style_only_sets_horizontal_alignment(self):
        styles = excel_style.bordered_styles(1, 1)["styles"]
        self.assertEqual(styles[excel_style.CENTER_STYLE_ID], {"align": "center"})

    def test_without_rows_no_cell_is_centered(self):
        cells = excel_style.bordered_styles(2, 3)["cell_styles"]
        self.assertTrue(all(excel_style.CENTER_STYLE_ID not in ids for ids in cells.values()))

    def test_header_row_height_has_a_30_point_minimum_and_grows_with_wrapped_headers(self):
        self.assertEqual(excel_style.header_row_height(["A", "B"], [12, 12]), 30)
        self.assertEqual(excel_style.header_row_height([], []), 30)
        self.assertGreater(excel_style.header_row_height(["Average Assessment Score Result Overall"], [10]), 30)
        # a header that fits on one or two lines never goes below the minimum
        self.assertEqual(excel_style.header_row_height(["Completion %"], [15]), 30)

    def test_make_xlsx_bytes_keeps_values_and_sets_header_height(self):
        import datetime

        rows = [["Name", "Score", "Date"], ["=1+1 text", 90, datetime.date(2026, 1, 2)]]
        styles = excel_style.bordered_styles(3, 2, rows=rows, column_formats={2: {"num_format": "d-mmm-yy"}})
        ws = openpyxl.load_workbook(BytesIO(excel_style.make_xlsx_bytes(rows, "Probe", styles))).active
        self.assertGreaterEqual(ws.row_dimensions[1].height, 30)
        self.assertEqual(ws["B2"].value, 90)
        self.assertEqual(ws["B2"].alignment.horizontal, "center")
        self.assertEqual(ws["C2"].alignment.horizontal, "center")
        self.assertEqual(ws["C2"].number_format, "d-mmm-yy")
        self.assertTrue(ws["A2"].alignment.wrap_text)


def _table_region(ws, first_header="Sr.No."):
    """(header_row, column_count) of the table, found by its first header
    ("Sr. No." on two exports, "Sr.No." on the Training Register, "Employee" on
    Employee Scores). Spaces are ignored when matching."""
    wanted = first_header.replace(" ", "")
    header_row = next(
        r for r in range(1, 10) if str(ws.cell(r, 1).value or "").replace(" ", "") == wanted
    )
    column_count = 0
    while ws.cell(header_row, column_count + 1).value not in (None, ""):
        column_count += 1
    return header_row, column_count


def _same(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return a == b


class TestDMSExcelExports(IntegrationTestCase):
    DOC = "DOC-2098-0001"

    def setUp(self):
        super().setUp()
        frappe.set_user("Administrator")
        # The transaction only rolls back when the class finishes, so the
        # synthetic fixtures are created once and shared by these read-only tests.
        if not frappe.db.exists("DMS Quiz", "QZ-XLS-A"):
            self._create_fixtures()

    def _create_fixtures(self):
        self._insert("DMS Quiz", name="QZ-XLS-A", title="QZ-XLS-A", pass_percentage=70)
        self._insert("Document Library", name=self.DOC, title="Excel Format Policy", training_quiz="QZ-XLS-A")
        for i, nm in enumerate(["Xena Test", "Yuri Test"], start=1):
            self._insert("Employee", name=f"TST-XLS-000{i}", employee_name=nm, department="Test Dept")
        self._insert("DMS Training Record", name="TRN-2098-0001", document=self.DOC, status="Assigned")
        self._insert(
            "Document Acknowledgement", parent="TRN-2098-0001", parenttype="DMS Training Record",
            parentfield="employees", employee="TST-XLS-0001", status="Completed", completion_date="2098-01-10",
        )
        self._insert(
            "Document Acknowledgement", parent="TRN-2098-0001", parenttype="DMS Training Record",
            parentfield="employees", employee="TST-XLS-0002", status="Pending",
        )
        for score, passed, when in ((50, 0, "2098-01-05 10:00:00"), (80, 1, "2098-01-10 09:00:00")):
            self._insert(
                "DMS Training Score History", employee="TST-XLS-0001", training_record="TRN-2098-0001",
                document=self.DOC, score=score, quiz_passed=passed, recorded_on=when,
            )

    @staticmethod
    def _insert(doctype, **fields):
        doc = frappe.new_doc(doctype)
        doc.update(fields)
        doc.db_insert()

    # ── helpers ───────────────────────────────────────────────────────────

    @staticmethod
    def _workbook_from_response():
        return openpyxl.load_workbook(BytesIO(frappe.response["filecontent"]))

    def _assert_table_formatting(self, ws, expected_headers):
        header_row, column_count = _table_region(ws, expected_headers[0])
        self.assertEqual(column_count, len(expected_headers))
        self.assertEqual([ws.cell(header_row, c).value for c in range(1, column_count + 1)], expected_headers)

        for c in range(1, column_count + 1):
            header = ws.cell(header_row, c)
            self.assertTrue(header.font.b, f"header {header.value!r} is not bold")
            self.assertEqual(header.alignment.horizontal, "center", f"header {header.value!r} not centered")
            self.assertTrue(header.alignment.wrap_text, f"header {header.value!r} not wrapped")

        # Every cell from the header down, in every table column -- including
        # blank ones -- needs a thin border on all four sides.
        for r in range(header_row, ws.max_row + 1):
            for c in range(1, column_count + 1):
                border = ws.cell(r, c).border
                for side in ("left", "right", "top", "bottom"):
                    self.assertEqual(
                        getattr(border, side).style, "thin", f"cell {ws.cell(r, c).coordinate} has no {side} border"
                    )

        for c in range(1, column_count + 1):
            letter = openpyxl.utils.get_column_letter(c)
            self.assertGreaterEqual(ws.column_dimensions[letter].width or 0, 8, f"column {letter} too narrow")

        # Wrapping on every table cell; numbers and real dates centered; text
        # (IDs, numeric-looking text, names) not centered.
        for r in range(header_row, ws.max_row + 1):
            for c in range(1, column_count + 1):
                cell = ws.cell(r, c)
                self.assertTrue(cell.alignment.wrap_text, f"{cell.coordinate} is not wrapped")
                if r == header_row:
                    continue
                if cell.value in (None, ""):
                    continue  # blank cells are bordered and wrapped, but have no value to align
                if cell.data_type in ("n", "d"):  # numbers and real dates
                    self.assertEqual(
                        cell.alignment.horizontal, "center", f"{cell.coordinate} ({cell.value!r}) is not centered"
                    )
                elif cell.value not in (None, ""):
                    self.assertNotEqual(
                        cell.alignment.horizontal, "center", f"text {cell.coordinate} ({cell.value!r}) is centered"
                    )

        self.assertGreaterEqual(ws.row_dimensions[header_row].height or 0, 30, "header row shorter than 30 points")
        return header_row, column_count

    def _body(self, ws, header_row, column_count):
        return [
            [ws.cell(r, c).value for c in range(1, column_count + 1)]
            for r in range(header_row + 1, ws.max_row + 1)
        ]

    def _assert_values(self, actual_rows, expected_rows):
        self.assertEqual(len(actual_rows), len(expected_rows))
        for i, (actual, expected) in enumerate(zip(actual_rows, expected_rows)):
            for j, (a, e) in enumerate(zip(actual, expected)):
                self.assertTrue(_same(a, e), f"row {i + 1} col {j + 1}: exported {a!r} != expected {e!r}")

    # ── 1. Training Analytics: Assessment Scores ──────────────────────────

    def test_assessment_scores_export_is_fully_bordered_with_bold_headers(self):
        api.export_assessment_scores()
        ws = self._workbook_from_response().active
        self.assertEqual(ws.title, "Assessment Scores")
        self._assert_table_formatting(ws, ASSESSMENT_SCORE_HEADERS)

    def test_assessment_scores_export_keeps_rows_values_and_order(self):
        scores = api._get_assessment_scores(None, None, None, None, None)
        expected = [
            [
                r["sr_no"], r["assessment_id"], r["training_title"] or "", r["date_of_completion"] or "",
                r["employee_id"], r["employee_name"] or "", r["department"] or "", r["trainer"] or "",
                r["maximum_marks"], r["marks_obtained"], r["score_pct"], r["pass_mark_pct"],
                r["result"], r["remarks"] or "",
            ]
            for r in scores
        ]
        self.assertTrue(any(row[2] == "Excel Format Policy" for row in expected), "fixture attempts missing")

        api.export_assessment_scores()
        ws = self._workbook_from_response().active
        header_row, column_count = _table_region(ws)
        self._assert_values(self._body(ws, header_row, column_count), expected)

    # ── 2. Training Register ──────────────────────────────────────────────

    def test_training_register_export_is_fully_bordered_with_bold_headers(self):
        api.export_training_register()
        ws = self._workbook_from_response().active
        expected = training_register.get_export_rows({})
        self._assert_table_formatting(ws, expected[0])

    def test_training_register_export_keeps_rows_values_and_order(self):
        expected = training_register.get_export_rows({})
        api.export_training_register()
        ws = self._workbook_from_response().active
        header_row, column_count = _table_region(ws)
        body = self._body(ws, header_row, column_count)
        self.assertGreater(len(body), 0)
        # Dates come back as datetimes in the workbook and as dates in the rows.
        normalised = [
            [v.date() if hasattr(v, "date") else v for v in row] for row in body
        ]
        self._assert_values(normalised, expected[1:])

    # ── 3. Training Assessment Results Register ───────────────────────────

    def test_assessment_register_export_table_is_fully_bordered_banner_excluded(self):
        api.export_training_assessment_results_register(training_document=self.DOC)
        ws = self._workbook_from_response().active
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertIn("A1:N1", merged)
        self.assertIn("A2:N2", merged)
        labels = [label for label, *_ in assessment_register.COLUMNS]
        header_row, _ = self._assert_table_formatting(ws, labels)
        self.assertGreater(header_row, 2, "banner rows must sit above the table")

    def test_assessment_register_export_keeps_rows_values_and_order(self):
        filters = {"training_document": self.DOC}
        expected_rows = assessment_register.get_register_rows(filters)
        fieldnames = [f for _l, f, *_ in assessment_register.COLUMNS]
        expected = [[row.get(f) for f in fieldnames] for row in expected_rows]
        self.assertEqual(len(expected), 3)  # 2 attempts + 1 pending employee

        api.export_training_assessment_results_register(training_document=self.DOC)
        ws = self._workbook_from_response().active
        header_row, column_count = _table_region(ws)
        body = [[v.date() if hasattr(v, "date") else v for v in row] for row in self._body(ws, header_row, column_count)]
        self._assert_values(body, expected)

    def test_assessment_register_all_documents_sheets_are_all_bordered(self):
        api.export_training_assessment_results_register(training_document=assessment_register.ALL_DOCUMENTS)
        wb = self._workbook_from_response()
        labels = [label for label, *_ in assessment_register.COLUMNS]
        self.assertGreaterEqual(len(wb.sheetnames), 2)
        for ws in wb.worksheets:
            self._assert_table_formatting(ws, labels)


# ── Safe workbook writer (used by the Employee Scores / Pending Trainings exports)

HOSTILE_TEXT = [
    "=1+1",
    "+1+1",
    "-2+3",
    "@SUM(1,1)",
    '=HYPERLINK("http://evil.example","click")',
    "http://example.com",
    "plain text",
    "007",
]


class TestSortKeyMatchesBrowserOrdering(UnitTestCase):
    """The Excel sort has to put names in the order the page does. The expected
    order below is what the browser's String.localeCompare (ICU, en-US) returns
    for the same names, captured from Node."""

    BROWSER_ORDER = [
        "_x", "#tag", "=1+1", "10abc", "1abc", "2abc", "a b", "a_b", "a-b", "ab", "apple", "Apple",
        "APPLE", "O'Brien", "OBrien", "Oersted", "Ørsted", "resume", "résumé", "strasse", "straße",
        "Xena Test", "Yuri Test", "Zed Plain",
    ]

    def test_server_sort_key_orders_names_like_the_browser(self):
        import random

        shuffled = self.BROWSER_ORDER[:]
        random.Random(7).shuffle(shuffled)
        self.assertEqual(sorted(shuffled, key=api._locale_sort_key), self.BROWSER_ORDER)

    def test_descending_swaps_operands_so_ties_keep_server_order(self):
        rows = [{"employee_name": "a", "n": 1}, {"employee_name": "A", "n": 2}, {"employee_name": "b", "n": 3}]
        spec = api._ANALYTICS_TABLES["employee_scores"]
        ordered = api._sort_analytics_rows(rows, spec, "completed", "desc")
        # `completed` is missing on every row -> -Infinity on both sides -> equal -> original order
        self.assertEqual([r["n"] for r in ordered], [1, 2, 3])

    def test_pending_comparator_treats_missing_as_empty_string(self):
        rows = [{"due_date_raw": "2097-02-01", "n": 1}, {"due_date_raw": None, "n": 2}, {"due_date_raw": "2097-01-01", "n": 3}]
        spec = api._ANALYTICS_TABLES["pending_trainings"]
        self.assertEqual([r["n"] for r in api._sort_analytics_rows(rows, spec, "due_date_raw", "asc")], [2, 3, 1])
        self.assertEqual([r["n"] for r in api._sort_analytics_rows(rows, spec, "due_date_raw", "desc")], [1, 3, 2])

    def test_numeric_comparator_puts_missing_values_last_in_both_directions(self):
        rows = [{"overall_score": None, "n": 1}, {"overall_score": 70.0, "n": 2}, {"overall_score": 90.0, "n": 3}]
        spec = api._ANALYTICS_TABLES["employee_scores"]
        self.assertEqual([r["n"] for r in api._sort_analytics_rows(rows, spec, "overall_score", "asc")], [2, 3, 1])
        self.assertEqual([r["n"] for r in api._sort_analytics_rows(rows, spec, "overall_score", "desc")], [3, 2, 1])


class TestSafeWorkbookBuilder(UnitTestCase):
    @staticmethod
    def _sheet(rows):
        return openpyxl.load_workbook(BytesIO(excel_style.build_workbook("Probe", rows))).active

    def test_text_that_looks_like_a_formula_or_link_stays_plain_text(self):
        ws = self._sheet([["Text"]] + [[text] for text in HOSTILE_TEXT])
        for row, text in enumerate(HOSTILE_TEXT, start=2):
            cell = ws.cell(row, 1)
            self.assertEqual(cell.data_type, "s", f"{text!r} was not written as text")
            self.assertEqual(cell.value, text)
            self.assertIsNone(cell.hyperlink, f"{text!r} became a hyperlink")

    def test_numbers_stay_numeric_and_blank_cells_stay_blank(self):
        ws = self._sheet([["N", "B"], [5, None], [2.5, ""], [0, "x"], [-1.5, None]])
        self.assertEqual([ws.cell(r, 1).data_type for r in range(2, 6)], ["n", "n", "n", "n"])
        self.assertEqual([ws.cell(r, 1).value for r in range(2, 6)], [5, 2.5, 0, -1.5])
        self.assertIsNone(ws.cell(2, 2).value)
        self.assertIsNone(ws.cell(3, 2).value)
        self.assertEqual(ws.cell(4, 2).value, "x")

    def test_every_cell_is_bordered_and_the_header_is_bold_centered_and_wrapped(self):
        ws = self._sheet([["A", "B"], [1, None], ["x", ""]])
        for r in range(1, 4):
            for c in range(1, 3):
                border = ws.cell(r, c).border
                for side in ("left", "right", "top", "bottom"):
                    self.assertEqual(getattr(border, side).style, "thin", f"{ws.cell(r, c).coordinate} {side}")
        for c in (1, 2):
            header = ws.cell(1, c)
            self.assertTrue(header.font.b)
            self.assertEqual(header.alignment.horizontal, "center")
            self.assertTrue(header.alignment.wrap_text)

    def test_numbers_and_real_dates_are_centered_and_text_is_not(self):
        import datetime

        ws = self._sheet(
            [["T", "N", "D"], ["HR-EMP-00013", 5, datetime.date(2026, 1, 2)], ["007", 2.5, datetime.datetime(2026, 1, 3, 9, 0)]]
        )
        for row in (2, 3):
            self.assertEqual(ws.cell(row, 2).alignment.horizontal, "center")
            self.assertEqual(ws.cell(row, 3).alignment.horizontal, "center")
            self.assertNotEqual(ws.cell(row, 1).alignment.horizontal, "center")
        self.assertEqual(ws["A3"].data_type, "s")  # numeric-looking text stays text
        self.assertEqual(ws["B2"].data_type, "n")

    def test_every_table_cell_wraps_text(self):
        ws = self._sheet([["A", "B", "C"], ["a long piece of text", 5, None], ["", 2.5, "x"]])
        for r in range(1, 4):
            for c in range(1, 4):
                self.assertTrue(ws.cell(r, c).alignment.wrap_text, ws.cell(r, c).coordinate)

    def test_header_row_height_is_at_least_30_points(self):
        ws = self._sheet([["A", "B"], [1, 2]])
        self.assertGreaterEqual(ws.row_dimensions[1].height, 30)
        # a narrow column forces a long header onto several lines -> taller row
        rows = [["Average Assessment Score Result Overall", "B"], [1, 2]]
        narrow = openpyxl.load_workbook(BytesIO(excel_style.build_workbook("Probe", rows, widths=[12, 10]))).active
        self.assertGreater(narrow.row_dimensions[1].height, 30)

    def test_borders_stop_at_the_table(self):
        ws = self._sheet([["A", "B"], [1, 2]])
        self.assertIsNone(ws.cell(3, 1).border.left.style)
        self.assertIsNone(ws.cell(1, 3).border.left.style)

    def test_column_widths_are_readable(self):
        ws = self._sheet([["Employee", "Department"], ["A very long employee name indeed", "HR"]])
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 20)
        self.assertGreaterEqual(ws.column_dimensions["B"].width, excel_style.MIN_WIDTH)

    def test_sheet_name_is_sanitised(self):
        wb = openpyxl.load_workbook(BytesIO(excel_style.build_workbook("A/B:C", [["x"]])))
        self.assertNotRegex(wb.sheetnames[0], r"[\[\]:*?/\\]")


# ── Employee Scores / Pending Trainings exports ─────────────────────────────

EMPLOYEE_HEADERS = ["Employee", "Department", "Completed", "Total", "Completion %", "Overall Score"]
PENDING_HEADERS = ["Employee", "Department", "Document", "Status", "Due Date"]


class TestAnalyticsTableExports(IntegrationTestCase):
    """Employee Scores and Pending Trainings Excel exports.

    Fixture cast (all in "Test Dept A/B", years 2097):
      Xena Test   A  completed 90 (due 01-10) + pending record 2 (due 02-01)
      Yuri Test   A  completed 70 (due 01-20)
      =1+1        B  overdue, no score (due 03-01)       <- formula-shaped name
      Zed Plain   B  pending, no due date                <- missing values
    """

    DOC_1 = "DOC-2097-0001"
    DOC_2 = "DOC-2097-0002"

    # reuse the formatting helpers of the export tests above
    _assert_table_formatting = TestDMSExcelExports._assert_table_formatting
    _body = TestDMSExcelExports._body
    _assert_values = TestDMSExcelExports._assert_values

    def setUp(self):
        super().setUp()
        frappe.set_user("Administrator")
        if not frappe.db.exists("Employee", "TST-ANL-0001"):
            self._create_fixtures()

    def tearDown(self):
        frappe.set_user("Administrator")
        super().tearDown()

    def _create_fixtures(self):
        for i, (name, dept) in enumerate(
            [("Xena Test", "Test Dept A"), ("Yuri Test", "Test Dept A"), ("=1+1", "Test Dept B"), ("Zed Plain", "Test Dept B")],
            start=1,
        ):
            self._insert("Employee", name=f"TST-ANL-000{i}", employee_name=name, department=dept)
        self._insert("DMS Training Record", name="TRN-2097-0001", document=self.DOC_1, status="Assigned")
        self._insert("DMS Training Record", name="TRN-2097-0002", document=self.DOC_2, status="Assigned")
        acks = [
            ("TRN-2097-0001", "TST-ANL-0001", "Completed", 90, "2097-01-10"),
            ("TRN-2097-0001", "TST-ANL-0002", "Completed", 70, "2097-01-20"),
            ("TRN-2097-0001", "TST-ANL-0003", "Overdue", 0, "2097-03-01"),
            ("TRN-2097-0002", "TST-ANL-0001", "Pending", 0, "2097-02-01"),
            ("TRN-2097-0002", "TST-ANL-0004", "Pending", 0, None),
        ]
        for parent, employee, status, score, due in acks:
            self._insert(
                "Document Acknowledgement", parent=parent, parenttype="DMS Training Record",
                parentfield="employees", employee=employee, status=status, assessment_score=score, due_date=due,
            )

    @staticmethod
    def _insert(doctype, **fields):
        doc = frappe.new_doc(doctype)
        doc.update(fields)
        doc.db_insert()

    # ── helpers ───────────────────────────────────────────────────────────

    def _export(self, table, **params):
        # Default: narrow the export to the fixtures with the department typeahead.
        params.setdefault("department_filter", "test dept")
        api.export_analytics_table(table=table, **params)
        ws = openpyxl.load_workbook(BytesIO(frappe.response["filecontent"])).active
        headers = EMPLOYEE_HEADERS if table == "employee_scores" else PENDING_HEADERS
        header_row, column_count = _table_region(ws, headers[0])
        return ws, self._body(ws, header_row, column_count)

    @staticmethod
    def _names(body):
        return [row[0] for row in body]

    @staticmethod
    def _due(date_string):
        return frappe.utils.format_date(date_string)

    XENA = ["Xena Test", "Test Dept A", 1, 2, 50.0, 90.0]
    YURI = ["Yuri Test", "Test Dept A", 1, 1, 100.0, 70.0]
    FORMULA = ["=1+1", "Test Dept B", 0, 1, 0.0, None]
    ZED = ["Zed Plain", "Test Dept B", 0, 1, 0.0, None]

    # ── schema + formatting ───────────────────────────────────────────────

    def test_employee_scores_schema_and_formatting(self):
        ws, _body = self._export("employee_scores")
        self.assertEqual(ws.title, "Employee Scores")
        self._assert_table_formatting(ws, EMPLOYEE_HEADERS)

    def test_pending_trainings_schema_and_formatting(self):
        ws, _body = self._export("pending_trainings")
        self.assertEqual(ws.title, "Pending Trainings")
        self._assert_table_formatting(ws, PENDING_HEADERS)

    def test_no_sr_no_column_and_csv_column_order(self):
        ws, _ = self._export("employee_scores")
        self.assertEqual([ws.cell(1, c).value for c in range(1, 8)], EMPLOYEE_HEADERS + [None])
        ws, _ = self._export("pending_trainings")
        self.assertEqual([ws.cell(1, c).value for c in range(1, 7)], PENDING_HEADERS + [None])

    # ── exported values ───────────────────────────────────────────────────

    def test_employee_scores_values_are_aggregated_correctly(self):
        _ws, body = self._export("employee_scores")
        by_name = {row[0]: row for row in body}
        self.assertEqual(set(by_name), {"Xena Test", "Yuri Test", "=1+1", "Zed Plain"})
        for expected in (self.XENA, self.YURI, self.FORMULA, self.ZED):
            self._assert_values([by_name[expected[0]]], [expected])

    def test_pending_trainings_values(self):
        _ws, body = self._export("pending_trainings")
        by_name = {row[0]: row for row in body}
        self.assertEqual(set(by_name), {"Xena Test", "=1+1", "Zed Plain"})
        self._assert_values(
            [by_name["Xena Test"], by_name["=1+1"], by_name["Zed Plain"]],
            [
                ["Xena Test", "Test Dept A", self.DOC_2, "Pending", self._due("2097-02-01")],
                ["=1+1", "Test Dept B", self.DOC_1, "Overdue", self._due("2097-03-01")],
                ["Zed Plain", "Test Dept B", self.DOC_2, "Pending", None],
            ],
        )

    def test_unfiltered_exports_equal_the_server_data_in_server_order(self):
        data = api.get_training_analytics()

        _ws, body = self._export("employee_scores", department_filter="")
        expected = [
            [r.get("employee_name") or r.get("employee"), r["department"], r["completed"], r["total"],
             r["completion_pct"], r["overall_score"]]
            for r in data["employee_scores"]
        ]
        self._assert_values(body, expected)

        _ws, body = self._export("pending_trainings", department_filter="")
        expected = [
            [r["employee_name"], r["department"], r["document"], r["status"], r["due_date"]]
            for r in data["pending_trainings"]
        ]
        self._assert_values(body, expected)

    def test_missing_values_are_blank_cells_and_numbers_stay_numeric(self):
        ws, _ = self._export("employee_scores", employee_filter="=1+1")
        self.assertIsNone(ws.cell(2, 6).value)  # no overall score
        for col in (3, 4, 5):
            self.assertEqual(ws.cell(2, col).data_type, "n")

        ws, _ = self._export("pending_trainings", employee_filter="zed")
        self.assertIsNone(ws.cell(2, 5).value)  # no due date

    # ── typeahead + page filters ──────────────────────────────────────────

    def test_employee_typeahead_is_case_insensitive(self):
        _ws, body = self._export("employee_scores", employee_filter="XENA")
        self.assertEqual(self._names(body), ["Xena Test"])
        _ws, body = self._export("pending_trainings", employee_filter="XENA")
        self.assertEqual(self._names(body), ["Xena Test"])

    def test_department_typeahead_is_a_contains_match(self):
        _ws, body = self._export("employee_scores", department_filter="dept b")
        self.assertEqual(sorted(self._names(body)), ["=1+1", "Zed Plain"])
        _ws, body = self._export("pending_trainings", department_filter="dept b")
        self.assertEqual(sorted(self._names(body)), ["=1+1", "Zed Plain"])

    def test_department_page_filter(self):
        _ws, body = self._export("employee_scores", department="Test Dept A", department_filter="")
        self.assertEqual(sorted(self._names(body)), ["Xena Test", "Yuri Test"])
        _ws, body = self._export("pending_trainings", department="Test Dept A", department_filter="")
        self.assertEqual(self._names(body), ["Xena Test"])

    def test_date_filters_follow_the_page_filters(self):
        _ws, body = self._export("employee_scores", from_date="2097-02-15")
        self.assertEqual(sorted(self._names(body)), ["=1+1", "Zed Plain"])

        _ws, body = self._export("employee_scores", to_date="2097-01-15")
        by_name = {row[0]: row for row in body}
        self.assertEqual(set(by_name), {"Xena Test", "Zed Plain"})
        # only the 01-10 assignment is in range for Xena now
        self._assert_values([by_name["Xena Test"]], [["Xena Test", "Test Dept A", 1, 1, 100.0, 90.0]])

        _ws, body = self._export("pending_trainings", to_date="2097-02-15")
        self.assertEqual(sorted(self._names(body)), ["Xena Test", "Zed Plain"])
        _ws, body = self._export("pending_trainings", from_date="2097-02-15")
        self.assertEqual(sorted(self._names(body)), ["=1+1", "Zed Plain"])

    # ── sorting ───────────────────────────────────────────────────────────

    def test_employee_scores_default_sort_is_overall_score_descending_missing_last(self):
        _ws, body = self._export("employee_scores")
        self.assertEqual(self._names(body)[:2], ["Xena Test", "Yuri Test"])
        self.assertEqual(sorted(self._names(body)[2:]), ["=1+1", "Zed Plain"])

    def test_employee_scores_missing_values_sort_last_in_both_directions(self):
        for direction, first_two in (("desc", ["Xena Test", "Yuri Test"]), ("asc", ["Yuri Test", "Xena Test"])):
            _ws, body = self._export("employee_scores", sort_key="overall_score", sort_dir=direction)
            self.assertEqual(self._names(body)[:2], first_two, direction)
            self.assertEqual(sorted(self._names(body)[2:]), ["=1+1", "Zed Plain"], direction)

    def test_employee_scores_sort_by_name_and_other_columns(self):
        _ws, body = self._export("employee_scores", sort_key="employee_name", sort_dir="asc")
        self.assertEqual(self._names(body), ["=1+1", "Xena Test", "Yuri Test", "Zed Plain"])
        _ws, body = self._export("employee_scores", sort_key="employee_name", sort_dir="desc")
        self.assertEqual(self._names(body), ["Zed Plain", "Yuri Test", "Xena Test", "=1+1"])

        _ws, body = self._export("employee_scores", sort_key="completion_pct", sort_dir="desc")
        self.assertEqual(self._names(body)[:2], ["Yuri Test", "Xena Test"])

        _ws, body = self._export("employee_scores", sort_key="department", sort_dir="asc")
        self.assertEqual(sorted(self._names(body)[:2]), ["Xena Test", "Yuri Test"])

        _ws, body = self._export("employee_scores", sort_key="completed", sort_dir="desc")
        self.assertEqual(self._names(body)[:2], ["Xena Test", "Yuri Test"])  # 1,1 (ties keep server order)

    def test_pending_trainings_default_sort_is_due_date_ascending_blank_first(self):
        # The page sorts with `a[key] || ""`, so a missing due date sorts as "".
        _ws, body = self._export("pending_trainings")
        self.assertEqual(self._names(body), ["Zed Plain", "Xena Test", "=1+1"])

    def test_pending_trainings_selected_column_sorts(self):
        _ws, body = self._export("pending_trainings", sort_key="due_date_raw", sort_dir="desc")
        self.assertEqual(self._names(body), ["=1+1", "Xena Test", "Zed Plain"])

        _ws, body = self._export("pending_trainings", sort_key="status", sort_dir="asc")
        self.assertEqual(self._names(body)[0], "=1+1")  # Overdue < Pending

        _ws, body = self._export("pending_trainings", sort_key="document", sort_dir="asc")
        self.assertEqual(self._names(body)[0], "=1+1")  # DOC-2097-0001 first

        _ws, body = self._export("pending_trainings", sort_key="employee_name", sort_dir="asc")
        self.assertEqual(self._names(body), ["=1+1", "Xena Test", "Zed Plain"])

    # ── formula safety ────────────────────────────────────────────────────

    def test_formula_shaped_text_is_exported_as_text_in_both_tables(self):
        for table in ("employee_scores", "pending_trainings"):
            ws, body = self._export(table, employee_filter="=1+1")
            self.assertEqual(len(body), 1, table)
            cell = ws.cell(2, 1)
            self.assertEqual(cell.value, "=1+1")
            self.assertEqual(cell.data_type, "s", f"{table}: employee name became a formula")

    # ── authorization + validation ────────────────────────────────────────

    def test_non_admin_users_are_rejected(self):
        for user in ("Guest",):
            frappe.set_user(user)
            for table in ("employee_scores", "pending_trainings"):
                with self.assertRaises(frappe.PermissionError):
                    api.export_analytics_table(table=table)

    def test_unknown_table_is_rejected(self):
        for bad in (None, "", "assessment_scores", "tabUser", "employee_scores; DROP TABLE x"):
            with self.assertRaises(frappe.ValidationError, msg=repr(bad)):
                api.export_analytics_table(table=bad)

    def test_invalid_sort_key_or_direction_is_rejected(self):
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", sort_key="bogus", sort_dir="asc")
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", sort_key="total", sort_dir="asc")  # not sortable on screen
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="pending_trainings", sort_key="overall_score", sort_dir="asc")
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", sort_key="overall_score", sort_dir="sideways")

    def test_invalid_dates_and_oversized_filters_are_rejected(self):
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", from_date="not-a-date")
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="pending_trainings", to_date="2097-99-99")
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", employee_filter="x" * 201)
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", department_filter="x" * 201)
        with self.assertRaises(frappe.ValidationError):
            api.export_analytics_table(table="employee_scores", department="x" * 201)

    # ── page wiring ───────────────────────────────────────────────────────

    def test_new_excel_buttons_cannot_trigger_the_assessment_export_handler(self):
        path = frappe.get_app_path("quality_dms", "dms", "page", "training_analytics", "training_analytics.js")
        with open(path, encoding="utf-8") as f:
            source = f.read()

        buttons = re.findall(r'<button([^>]*)>', source)
        table_excel = [b for b in buttons if "ta-download-excel-table" in b]
        self.assertEqual(len(table_excel), 2)
        self.assertEqual(
            sorted(re.search(r'data-table="([^"]+)"', b).group(1) for b in table_excel),
            ["employee_scores", "pending_trainings"],
        )

        def class_tokens(attrs):
            match = re.search(r'class="([^"]*)"', attrs)
            return match.group(1).split() if match else []

        for attrs in table_excel:
            tokens = class_tokens(attrs)
            self.assertNotIn("ta-download-excel", tokens)  # the Assessment Scores handler's selector
            self.assertNotIn("ta-download", tokens)  # the CSV handler's selector

        # exactly one element still carries the Assessment Scores export class
        self.assertEqual(sum("ta-download-excel" in class_tokens(b) for b in buttons), 1)
